import csv
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=13, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_note = Font(name=F, size=9, italic=True, color="666666")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_hdr = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF")
USD = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.00%;(0.00%);-'
DT = 'yyyy-mm-dd'

wb = Workbook()

def hdr(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = f_bold; c.fill = fill_hdr
        c.alignment = Alignment(wrap_text=True, vertical="center")

def inp(c, v, fmt=None, key=False):
    c.value = v; c.font = f_in
    if key: c.fill = fill_in
    if fmt: c.number_format = fmt

def widths(ws, ws_w):
    for col, w in ws_w.items(): ws.column_dimensions[col].width = w

# ---------------- Guide ----------------
g = wb.active; g.title = "Guide"
lines = [
 ("NEST buyback model (Lido DAO), v1, 28.09.2026", f_title),
 ("", f_norm),
 ("How to use", f_bold),
 ("1. All buyback rule settings are on the Params sheet, in yellow cells with blue text. Change a value and everything recalculates.", f_norm),
 ("2. The result for the selected settings is on the Summary sheet: buybacks by year, days with purchases, budget at period end.", f_norm),
 ("3. Future scenarios via four levers (TVL, yield, DAO share, ETH price) are on the Forward sheet, also in yellow cells.", f_norm),
 ("", f_norm),
 ("Sheets", f_bold),
 ("Data: 1,002 actual fee inflows to the DAO treasury since 01.01.2024, one per daily rebase. Transaction hash, stETH, Chainlink prices (roundId). USD is computed by formula.", f_norm),
 ("Params: buyback rule parameters.", f_norm),
 ("Expenses: Foundations expenses by period from Lido reports and cost of revenue (referral payouts and other).", f_norm),
 ("Sim: daily simulation of the NEST budget using contract formulas (threshold, share, caps, minimum purchase, deficit policy).", f_norm),
 ("NEST_check: reproduction of the NEST onchain budget at all checkpoints. Error should be 0.", f_norm),
 ("Recon: true-up of onchain inflows against Lido reports.", f_norm),
 ("Forward: scenarios 12 months ahead.", f_norm),
 ("Grid_Python: scenario grid from the Python model, for true-up (values, not formulas).", f_norm),
 ("", f_norm),
 ("Colors", f_bold),
 ("Blue text on yellow background: inputs you can change. Blue text without background: source data (onchain or from reports). Black: formulas.", f_norm),
 ("", f_norm),
 ("What the revenue series is", f_bold),
 ("Actual staking fee inflows to the treasury (after the node operator share, before referral payouts), valued at Chainlink stETH/ETH × ETH/USD at the inflow block. This is the same oracle route as NEST uses, but not its historical counter: NEST converts to USD later, in a separate call. The stETH amount matches NEST exactly, USD differs by −0.044%.", f_norm),
 ("Methodology details and all checks are in base/README.md.", f_norm),
 ("", f_norm),
 ("Assumptions", f_bold),
 ("Expenses for 2024, 2025 and H2 2026 by line item (staking, shared services) are unknown; for threshold variants 2 and 3 they are split using H1 2026 proportions. H2 2026 expenses = annual forecast $37.7M minus H1 actuals.", f_norm),
 ("The annual cap is computed in 365-day windows from the simulation start date, as the contract does from the activation date.", f_norm),
]
for i, (t, f) in enumerate(lines, 1):
    c = g.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
g.column_dimensions["A"].width = 130

# ---------------- Data ----------------
d = wb.create_sheet("Data")
rows = list(csv.DictReader(open("base/treasury_fee_inflows.csv")))
hdr(d, 1, ["Date (UTC)", "stETH to treasury", "stETH/ETH (Chainlink)", "ETH/USD (Chainlink)", "USD value", "Source",
           "Block", "Tx hash", "stETH/ETH roundId", "ETH/USD roundId", "Price flag"])
for i, r in enumerate(rows, 2):
    d.cell(row=i, column=1, value=date.fromisoformat(r["date"])).number_format = DT
    inp(d.cell(row=i, column=2), float(r["steth"]), "#,##0.000000")
    inp(d.cell(row=i, column=3), float(r["steth_eth"]), "0.000000")
    inp(d.cell(row=i, column=4), float(r["eth_usd"]), "#,##0.00")
    d.cell(row=i, column=5, value=f"=B{i}*C{i}*D{i}").number_format = USD2
    d.cell(row=i, column=6, value=r["source"])
    d.cell(row=i, column=7, value=int(r["block"]))
    d.cell(row=i, column=8, value=r["tx_hash"])
    d.cell(row=i, column=9, value=r["steth_eth_roundId"])
    d.cell(row=i, column=10, value=r["eth_usd_roundId"])
    d.cell(row=i, column=11, value=r["price_flag"])
    for col in (1, 5, 6, 7, 8, 9, 10, 11): d.cell(row=i, column=col).font = f_norm
N = len(rows); LAST = N + 1
d.freeze_panes = "A2"
widths(d, {"A": 12, "B": 16, "C": 14, "D": 14, "E": 14, "F": 24, "G": 11, "H": 68, "I": 22, "J": 22, "K": 26})
d.cell(row=1, column=2).comment = Comment("Actual stETH transfer to the treasury agent 0x3e40...9C8c: until 24.12.2025 minted at rebase, afterwards transferred from distributor 0x23ed...cddf. Source: base/treasury_fee_inflows.csv", "model")
d.cell(row=1, column=3).comment = Comment("Chainlink stETH/ETH proxy 0x86392dC19c0b719886221c78AB11eb8Cf5c52812, latestRoundData() at the inflow block", "model")
d.cell(row=1, column=4).comment = Comment("Chainlink ETH/USD proxy 0x5f4eC3Df9cbd43714FE2740f5E3616155c5b8419, latestRoundData() at the inflow block", "model")

DA = f"Data!$A$2:$A${LAST}"; DB = f"Data!$B$2:$B${LAST}"; DD = f"Data!$D$2:$D${LAST}"; DE = f"Data!$E$2:$E${LAST}"

# ---------------- Params ----------------
p = wb.create_sheet("Params")
p["A1"] = "Buyback rule parameters"; p["A1"].font = f_title
p["A2"] = "Change only the yellow cells. Codes are in column C."; p["A2"].font = f_note
hdr(p, 3, ["Parameter", "Value", "Explanation"])
params = [
 (4, "Simulation start", date(2025, 1, 1), DT, "From this date the budget starts at the initial value (B17)"),
 (5, "Simulation end", date(2026, 9, 28), DT, f"Data available through {rows[-1]['date']}"),
 (6, "Threshold type", 1, "0", "1 = fixed threshold (B8), 2 = threshold based on Foundations expenses (variant in B7)"),
 (7, "Expense variant for threshold", 1, "0", "1 = all Foundations expenses, 2 = staking + shared services, 3 = staking only"),
 (8, "Fixed threshold, $ per year", 40000000, USD, "NEST now: $40M (LIP-36)"),
 (9, "Expense multiplier", 1, "0.00", "1 = as in reports, 0.8 = expenses 20% lower"),
 (10, "Revenue base", 1, "0", "1 = treasury inflows (as NEST), 2 = net, after cost of revenue"),
 (11, "Surplus share to buybacks", 0.5, "0%", "NEST now: 50%"),
 (12, "Daily cap, $", 50000, USD, "NEST now: $50k. 0 = no cap"),
 (13, "Annual cap, $", 10000000, USD, "NEST now: $10M per 365-day window. 0 = no cap"),
 (14, "Minimum purchase, $", 1000, USD, "NEST now: $1,000 (minSpendPerCallUSD)"),
 (15, "Deficit policy", 1, "0", "1 = deficit accumulates without limit (as now), 2 = floor of N days (B16), 3 = deficit does not accumulate"),
 (16, "Deficit floor, days of threshold", 30, "0", "Used when B15 = 2"),
 (17, "Initial budget, $", 0, USD, "E.g. −548,925 = NEST budget as of 28.09.2026"),
]
for r, name, v, fmt, note in params:
    p.cell(row=r, column=1, value=name).font = f_norm
    inp(p.cell(row=r, column=2), v, fmt, key=True)
    c = p.cell(row=r, column=3, value=note); c.font = f_note
widths(p, {"A": 34, "B": 18, "C": 95})
for ref, lst in [("B6", '"1,2"'), ("B7", '"1,2,3"'), ("B10", '"1,2"'), ("B15", '"1,2,3"')]:
    dv = DataValidation(type="list", formula1=lst, allow_blank=False); p.add_data_validation(dv); dv.add(ref)
p["A19"] = "Presets for comparison (enter manually in B6..B17):"; p["A19"].font = f_bold
presets = ["Current NEST: B6=1, B8=40M, B10=1, B11=50%, B15=1",
           "Aksusarya proposal: B6=1, B8=30M, B11=100%, B15=1",
           "Expense-based threshold, fair: B6=2, B7=1, B10=2, B11=100%, B15=3",
           "Same with expenses −20%: plus B9=0.8"]
for i, t in enumerate(presets, 20):
    p.cell(row=i, column=1, value=t).font = f_norm

# ---------------- Expenses ----------------
e = wb.create_sheet("Expenses")
e["A1"] = "Foundations expenses and cost of revenue by period"; e["A1"].font = f_title
hdr(e, 4, ["Period", "Start", "End", "Days", "Foundations expenses, $", "Staking (core+growth), $", "Shared services, $",
           "Onchain inflows, $", "Cost of revenue, $", "Cost share", "Daily threshold: all expenses",
           "Daily threshold: staking+shared", "Daily threshold: staking", "Report: net staking revenue, $", "Source / assumption"])
per = [
 ("2024", date(2024,1,1), date(2024,12,31)),
 ("2025", date(2025,1,1), date(2025,12,31)),
 ("Q1 2026", date(2026,1,1), date(2026,3,31)),
 ("Q2 2026", date(2026,4,1), date(2026,6,30)),
 ("H2 2026", date(2026,7,1), date(2026,12,31)),
]
for i, (lab, s, en) in enumerate(per, 5):
    e.cell(row=i, column=1, value=lab).font = f_norm
    e.cell(row=i, column=2, value=s).number_format = DT
    e.cell(row=i, column=3, value=en).number_format = DT
    e.cell(row=i, column=4, value=f"=C{i}-B{i}+1")
    e.cell(row=i, column=8, value=f'=SUMIFS({DE},{DA},">="&B{i},{DA},"<="&C{i})').number_format = USD
    e.cell(row=i, column=10, value=f"=IF(H{i}>0,I{i}/H{i},0)").number_format = PCT
    e.cell(row=i, column=11, value=f"=E{i}/D{i}*1").number_format = USD
    e.cell(row=i, column=12, value=f"=(F{i}+G{i})/D{i}").number_format = USD
    e.cell(row=i, column=13, value=f"=F{i}/D{i}").number_format = USD
# expenses totals
inp(e["E5"], 52100000, USD); inp(e["E6"], 45500000, USD); inp(e["E7"], 6920000, USD); inp(e["E8"], 7410000, USD)
e["E9"] = "=Expenses!$B$13-E7-E8"; e["E9"].number_format = USD
inp(e["F7"], 4380000, USD); inp(e["F8"], 4600000, USD); inp(e["G7"], 1570000, USD); inp(e["G8"], 1580000, USD)
for r in (5, 6, 9):
    e[f"F{r}"] = f"=E{r}*$B$14"; e[f"G{r}"] = f"=E{r}*$B$15"
    e[f"F{r}"].number_format = USD; e[f"G{r}"].number_format = USD
# cost of revenue
inp(e["I5"], 3100000, USD); inp(e["I6"], 4100000, USD)
inp(e["N7"], 8630000, USD); inp(e["N8"], 7080000, USD)
e["I7"] = "=H7-N7"; e["I8"] = "=H8-N8"; e["I9"] = "=H9*$B$16"
for c in ("I7", "I8", "I9"): e[c].number_format = USD
e["O5"] = "Expenses and referral payouts: GOOSE-2025 & EGGs-2025 Final Report (research.lido.fi/t/11304)"
e["O6"] = "Same"
e["O7"] = "H1 2026 Report, P&L (lido.fi/ldo-hub/reports/h1-2026); costs = onchain − net revenue"
e["O8"] = "Same"
e["O9"] = "Expenses = annual forecast (B13) − H1 actuals; line items and cost share by H1 proportions (assumption)"
for r in range(5, 10): e[f"O{r}"].font = f_note
e["A12"] = "Inputs and derived values"; e["A12"].font = f_bold
e["A13"] = "Expense forecast for 2026, $"; inp(e["B13"], 37700000, USD, key=True)
e["C13"] = "H1 2026 Report: full-year spending projected at ~$37.7M"; e["C13"].font = f_note
e["A14"] = "Staking share of expenses (H1 2026)"; e["B14"] = "=(F7+F8)/(E7+E8)"; e["B14"].number_format = PCT
e["A15"] = "Shared services share of expenses (H1 2026)"; e["B15"] = "=(G7+G8)/(E7+E8)"; e["B15"].number_format = PCT
e["A16"] = "Cost of revenue share (H1 2026)"; e["B16"] = "=(I7+I8)/(H7+H8)"; e["B16"].number_format = PCT
e["A17"] = "Onchain H2 2026 is incomplete (data up to the date of the last Data row); the H2 cost share is taken from H1."; e["A17"].font = f_note
widths(e, {"A": 40, "B": 14, "C": 12, "D": 7, "E": 17, "F": 17, "G": 15, "H": 17, "I": 17, "J": 10, "K": 14, "L": 14, "M": 14, "N": 17, "O": 80})

# ---------------- Sim ----------------
s = wb.create_sheet("Sim")
hdr(s, 1, ["Date", "Active", "Period #", "Treasury inflow, $", "Revenue used, $", "Baseline/day, $", "Budget delta, $",
           "Budget before floor, $", "Budget after floor, $", "Cap window #", "Spent in window before, $", "Buy, $",
           "Budget end of day, $", "Year"])
for i in range(2, LAST + 1):
    pr = i - 1
    s[f"A{i}"] = f"=Data!A{i}"; s[f"A{i}"].number_format = DT
    s[f"B{i}"] = f"=IF(AND(A{i}>=Params!$B$4,A{i}<=Params!$B$5),1,0)"
    s[f"C{i}"] = f"=MATCH(A{i},Expenses!$B$5:$B$9,1)"
    s[f"D{i}"] = f"=Data!E{i}"
    s[f"E{i}"] = f"=IF(Params!$B$10=1,D{i},D{i}*(1-INDEX(Expenses!$J$5:$J$9,C{i})))"
    s[f"F{i}"] = f"=IF(Params!$B$6=1,Params!$B$8/365,INDEX(Expenses!$K$5:$M$9,C{i},Params!$B$7)*Params!$B$9)"
    s[f"G{i}"] = f"=IF(B{i}=1,Params!$B$11*(E{i}-F{i}),0)"
    prev = "Params!$B$17" if i == 2 else f"IF(B{pr}=1,M{pr},Params!$B$17)"
    s[f"H{i}"] = f"=IF(B{i}=1,{prev}+G{i},0)"
    s[f"I{i}"] = (f"=IF(B{i}=1,IF(Params!$B$15=3,MAX(0,H{i}),IF(Params!$B$15=2,"
                  f"MAX(H{i},-Params!$B$16*F{i}*Params!$B$11),H{i})),0)")
    s[f"J{i}"] = f"=IF(B{i}=1,INT((A{i}-Params!$B$4)/365),-1)"
    s[f"K{i}"] = "=0" if i == 2 else f"=IF(B{i}=1,SUMIFS(L$2:L{pr},J$2:J{pr},J{i}),0)"
    s[f"L{i}"] = (f"=IF(AND(B{i}=1,I{i}>Params!$B$14),MAX(0,MIN(I{i},IF(Params!$B$12>0,Params!$B$12,1E+18),"
                  f"IF(Params!$B$13>0,Params!$B$13-K{i},1E+18))),0)")
    s[f"M{i}"] = f"=I{i}-L{i}"
    s[f"N{i}"] = f"=YEAR(A{i})"
    for col in "DEFGHIKLM": s[f"{col}{i}"].number_format = USD
s.freeze_panes = "B2"
widths(s, {"A": 12, "B": 7, "C": 8, "D": 15, "E": 15, "F": 14, "G": 14, "H": 16, "I": 16, "J": 9, "K": 16, "L": 12, "M": 16, "N": 7})

SL = f"Sim!$L$2:$L${LAST}"; SN = f"Sim!$N$2:$N${LAST}"; SB = f"Sim!$B$2:$B${LAST}"; SE = f"Sim!$E$2:$E${LAST}"
SF = f"Sim!$F$2:$F${LAST}"; SA = f"Sim!$A$2:$A${LAST}"; SM = f"Sim!$M$2:$M${LAST}"

# ---------------- Summary ----------------
m = wb.create_sheet("Summary", 1)
m["A1"] = "Result for current parameters (Params sheet)"; m["A1"].font = f_title
hdr(m, 3, ["Year", "Revenue (used), $", "Threshold for period, $", "Buybacks, $", "Days with purchases"])
for i, y in enumerate([2024, 2025, 2026], 4):
    m[f"A{i}"] = str(y)
    m[f"B{i}"] = f"=SUMIFS({SE},{SN},{y},{SB},1)"
    m[f"C{i}"] = f"=SUMIFS({SF},{SN},{y},{SB},1)"
    m[f"D{i}"] = f"=SUMIFS({SL},{SN},{y})"
    m[f"E{i}"] = f'=COUNTIFS({SN},{y},{SL},">0")'
    for c in "BCD": m[f"{c}{i}"].number_format = USD
m["A7"] = "Total"; m["A7"].font = f_bold
for c in "BCDE":
    m[f"{c}7"] = f"=SUM({c}4:{c}6)"; m[f"{c}7"].font = f_bold
    if c != "E": m[f"{c}7"].number_format = USD
m["A9"] = "First purchase"; m["B9"] = f'=IF(E7>0,_xlfn.MINIFS({SA},{SL},">0"),"no")'; m["B9"].number_format = DT
m["A10"] = "Budget at period end, $"; m["B10"] = f"=INDEX({SM},MATCH(Params!$B$5,{SA},0))"; m["B10"].number_format = USD
m["A11"] = "Deepest deficit in period, $"; m["B11"] = f"=MIN({SM})"; m["B11"].number_format = USD
m["A13"] = "True-up against the Python model: with the ‘Current NEST’ preset from 01.01.2025, buybacks = $1,073,350, budget = −$2,833,920 (Grid_Python sheet)."; m["A13"].font = f_note
widths(m, {"A": 36, "B": 22, "C": 20, "D": 16, "E": 16})

# ---------------- NEST_check ----------------
n = wb.create_sheet("NEST_check")
n["A1"] = "Reproduction of the NEST onchain budget at checkpoints"; n["A1"].font = f_title
n["A2"] = "Contract formula: budget = share × (counter revenue − daily threshold × number of daily slots). Slots are counted from the activation day inclusive through the checkpoint day inclusive."; n["A2"].font = f_note
n["A4"] = "NEST activation (activationTS)"; inp(n["B4"], date(2026, 8, 10), DT)
n["A5"] = "Daily threshold, $ (reserveDailyRateUSD)"; inp(n["B5"], 109589, USD)
n["A6"] = "Share (surplusShareBP)"; inp(n["B6"], 0.5, "0%")
n["A7"] = "First rebase captured by the counter"; inp(n["B7"], date(2026, 8, 15), DT)
n["C4"] = "Source: BuybackAllocator 0xAA568141c051f2D1132b110f8391F18D48E8D889, StakingRevenueSource 0x6220212a33a87Ed7Cc386B67eB2c393974F28C38"; n["C4"].font = f_note
hdr(n, 9, ["Checkpoint date", "Block", "Tx hash", "NEST revenue counter, $", "NEST onchain budget, $", "Slots",
           "Budget by formula, $", "Error, $", "Our revenue for the same rebases, $", "Budget on our revenue, $", "Difference from onchain, %"])
cps = list(csv.DictReader(open("base/nest_checkpoints.csv")))
for i, c in enumerate(cps, 10):
    dt = date.fromisoformat(c["ts"][:10])
    n[f"A{i}"] = dt; n[f"A{i}"].number_format = DT
    n[f"B{i}"] = int(c["block"]); n[f"C{i}"] = c["tx"]
    inp(n[f"D{i}"], float(c["totalRevenueUSD"]), USD2); inp(n[f"E{i}"], float(c["budgetUSD"]), USD2)
    n[f"F{i}"] = f"=A{i}-$B$4+1"
    n[f"G{i}"] = f"=$B$6*(D{i}-$B$5*F{i})"; n[f"G{i}"].number_format = USD2
    n[f"H{i}"] = f"=G{i}-E{i}"; n[f"H{i}"].number_format = USD2
    n[f"I{i}"] = f'=SUMIFS({DE},{DA},">="&$B$7,{DA},"<"&A{i})'; n[f"I{i}"].number_format = USD2
    n[f"J{i}"] = f"=$B$6*(I{i}-$B$5*F{i})"; n[f"J{i}"].number_format = USD2
    n[f"K{i}"] = f"=IF(D{i}>0,I{i}/D{i}-1,0)"; n[f"K{i}"].number_format = PCT
lastn = 9 + len(cps)
n[f"A{lastn+2}"] = "Maximum formula error, $"; n[f"B{lastn+2}"] = f"=MAX(MAX(H10:H{lastn}),-MIN(H10:H{lastn}))"
n[f"B{lastn+2}"].number_format = USD2
n[f"A{lastn+3}"] = "Latest NEST budget, $"; n[f"B{lastn+3}"] = f"=E{lastn}"; n[f"B{lastn+3}"].number_format = USD2
widths(n, {"A": 36, "B": 12, "C": 68, "D": 20, "E": 20, "F": 8, "G": 20, "H": 12, "I": 22, "J": 22, "K": 12})
NEST_LAST_BUDGET = f"NEST_check!$B${lastn+3}"

# ---------------- Recon ----------------
rc = wb.create_sheet("Recon")
rc["A1"] = "True-up of onchain inflows against Lido reports"; rc["A1"].font = f_title
hdr(rc, 3, ["Period", "Start", "End", "Onchain treasury inflows, $", "Report: net staking revenue, $", "Difference, $",
            "Report: referral payouts, $", "Residual, $", "Report: total protocol fee, $", "Treasury share of fee", "Source"])
recon = [
 ("2024", date(2024,1,1), date(2024,12,31), 48500000, 3100000, None, "GOOSE-2025 Final Report (t/11304)"),
 ("2025", date(2025,1,1), date(2025,12,31), 37400000, 4100000, None, "GOOSE-2025 Final Report (t/11304)"),
 ("Q1 2026", date(2026,1,1), date(2026,3,31), 8630000, None, 14990000, "H1 2026 Report, P&L"),
 ("Q2 2026", date(2026,4,1), date(2026,6,30), 7080000, None, 12520000, "H1 2026 Report, P&L"),
 ("H1 2026", date(2026,1,1), date(2026,6,30), 15710000, None, 27510000, "H1 2026 Report, P&L"),
]
for i, (lab, a, b, net, ref, gross, src) in enumerate(recon, 4):
    rc[f"A{i}"] = lab; rc[f"B{i}"] = a; rc[f"C{i}"] = b
    rc[f"B{i}"].number_format = DT; rc[f"C{i}"].number_format = DT
    rc[f"D{i}"] = f'=SUMIFS({DE},{DA},">="&B{i},{DA},"<="&C{i})'
    inp(rc[f"E{i}"], net, USD)
    rc[f"F{i}"] = f"=D{i}-E{i}"
    if ref is not None:
        inp(rc[f"G{i}"], ref, USD); rc[f"H{i}"] = f"=F{i}-G{i}"
    else:
        rc[f"G{i}"] = "not disclosed"; rc[f"H{i}"] = "n/a"
    if gross is not None:
        inp(rc[f"I{i}"], gross, USD); rc[f"J{i}"] = f"=D{i}/I{i}"; rc[f"J{i}"].number_format = PCT
    rc[f"K{i}"] = src; rc[f"K{i}"].font = f_note
    for c in "DFH": rc[f"{c}{i}"].number_format = USD
rc["A10"] = "The 2024/2025 residual is within report rounding ($0.1M) and the difference in valuation methods (report at end-of-day price, onchain at block price)."; rc["A10"].font = f_note
rc["A11"] = "The H1 2026 difference is cost of revenue: under Lido's methodology, referral payouts and possibly operator discounts; there is no breakdown for H1."; rc["A11"].font = f_note
widths(rc, {"A": 10, "B": 12, "C": 12, "D": 22, "E": 22, "F": 14, "G": 18, "H": 12, "I": 20, "J": 12, "K": 36})

# ---------------- Forward ----------------
fw = wb.create_sheet("Forward")
fw["A1"] = "12-month scenarios: revenue = TVL × yield × DAO share × ETH price"; fw["A1"].font = f_title
fw["A3"] = "Calibration to actual data"; fw["A3"].font = f_bold
fw["A4"] = "Average treasury inflows over 30 days, stETH/day"
fw["B4"] = f'=AVERAGEIFS({DB},{DA},">"&(MAX({DA})-30))'; fw["B4"].number_format = "0.00"
fw["A5"] = "Latest ETH/USD price"; fw["B5"] = f"=INDEX({DD},MATCH(MAX({DA}),{DA},0))"; fw["B5"].number_format = "#,##0.00"
fw["A6"] = "DAO share of rewards"; inp(fw["B6"], 0.0617, PCT, key=True)
fw["C6"] = "StakingRouter getStakingFeeAggregateDistributionE4Precision(), 28.09.2026"; fw["C6"].font = f_note
fw["A7"] = "Lido TVL, M ETH"; inp(fw["B7"], 9.13, "0.00", key=True)
fw["C7"] = "H1 2026 Report, 30.06.2026, with entry queue; the implied yield below is calibrated so that revenue matches actuals"; fw["C7"].font = f_note
fw["A8"] = "Implied gross staking yield"; fw["B8"] = "=B4*365/(B6*B7*1000000)"; fw["B8"].number_format = PCT
fw["A9"] = "Current NEST budget, $"; fw["B9"] = f"={NEST_LAST_BUDGET}"; fw["B9"].number_format = USD; fw["B9"].font = f_link
fw["A10"] = "Annual cap, $"; inp(fw["B10"], 10000000, USD)
fw["A11"] = "Share for expense-based rule"; inp(fw["B11"], 1.0, "0%", key=True)

hdr(fw, 13, ["Metric", "Now", "Team targets (23% share)", "ETH $2,000", "ETH $3,500", "Issuance −30%", "Move to Curated v2 (5.8%)"])
cols = "BCDEFG"
fw["A14"] = "TVL, M ETH"; fw["A15"] = "Gross yield"; fw["A16"] = "DAO share"; fw["A17"] = "ETH price, $"
fw["A18"] = "Cost of revenue share"; fw["A19"] = "Foundations expenses, $ per year"; fw["A20"] = "NEST fixed threshold, $"
for c in cols:
    fw[f"{c}14"] = "=$B$7"; fw[f"{c}15"] = "=$B$8"; fw[f"{c}16"] = "=$B$6"; fw[f"{c}17"] = "=$B$5"
    fw[f"{c}18"] = "=Expenses!$B$16"; fw[f"{c}19"] = "=Expenses!$B$13"
    fw[f"{c}20"] = "=Params!$B$8"
fw["C14"] = "=$B$7*23/21.2"
fw["D17"] = 2000; fw["D17"].font = f_in; fw["D17"].fill = fill_in
fw["E17"] = 3500; fw["E17"].font = f_in; fw["E17"].fill = fill_in
fw["F15"] = "=$B$8*0.7"
fw["G16"] = 0.058; fw["G16"].font = f_in; fw["G16"].fill = fill_in
fw["A22"] = "Results"; fw["A22"].font = f_bold
labels = {23: "Treasury inflows, $ per year", 24: "Net staking revenue, $ per year",
          25: "Break-even ETH price for the NEST threshold", 26: "ETH price at which net revenue = expenses",
          27: "Buybacks under current rule (50%, cap), $ per year", 28: "Days to pay off current deficit (current rule)",
          29: "Buybacks: threshold = expenses, net revenue, share B11, $ per year", 30: "Buybacks: proposal 30M / 100%, $ per year"}
for r, t in labels.items(): fw[f"A{r}"] = t
for c in cols:
    fw[f"{c}23"] = f"={c}14*1000000*{c}15*{c}16*{c}17"
    fw[f"{c}24"] = f"={c}23*(1-{c}18)"
    fw[f"{c}25"] = f"={c}20/({c}14*1000000*{c}15*{c}16)"
    fw[f"{c}26"] = f"={c}19/({c}14*1000000*{c}15*{c}16*(1-{c}18))"
    fw[f"{c}27"] = f"=MIN($B$10,0.5*MAX(0,{c}23-{c}20))"
    fw[f"{c}28"] = f'=IF({c}23>{c}20,-$B$9/(0.5*({c}23-{c}20)/365),"not paid off")'
    fw[f"{c}29"] = f"=MIN($B$10,$B$11*MAX(0,{c}24-{c}19))"
    fw[f"{c}30"] = f"=MIN($B$10,MAX(0,{c}23-30000000))"
    for r in (19, 20, 23, 24, 27, 29, 30): fw[f"{c}{r}"].number_format = USD
    for r in (25, 26, 17): fw[f"{c}{r}"].number_format = "$#,##0"
    fw[f"{c}14"].number_format = "0.00"; fw[f"{c}15"].number_format = PCT; fw[f"{c}16"].number_format = PCT; fw[f"{c}18"].number_format = PCT
    fw[f"{c}28"].number_format = "#,##0"
fw["A32"] = "Simplifications: flat revenue for the year ahead, the daily cap does not bind ($50k × 365 > $10M), the deficit is paid off only in row 28."; fw["A32"].font = f_note
widths(fw, {"A": 58, "B": 16, "C": 18, "D": 14, "E": 14, "F": 14, "G": 20})

# ---------------- Grid_Python ----------------
gp = wb.create_sheet("Grid_Python")
gp["A1"] = "Scenario grid from the Python model (simulator.py) on the same base, ending 28.09.2026. Values, not formulas."; gp["A1"].font = f_bold
grid = list(csv.reader(open("results_grid_base.csv")))
hdr(gp, 3, grid[0])
for i, r in enumerate(grid[1:], 4):
    for j, v in enumerate(r, 1):
        try: v = float(v) if j >= 3 and j <= 8 else v
        except ValueError: pass
        c = gp.cell(row=i, column=j, value=v); c.font = f_norm
        if 3 <= j <= 7: c.number_format = USD
widths(gp, {"A": 13, "B": 32, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 12, "I": 12})

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != F:
                c.font = Font(name=F, size=10, bold=bool(c.font.bold), italic=bool(c.font.italic), color=(c.font.color.rgb if c.font.color is not None and isinstance(c.font.color.rgb, str) else None))

wb.save("../nest_model.xlsx")
print("saved")
