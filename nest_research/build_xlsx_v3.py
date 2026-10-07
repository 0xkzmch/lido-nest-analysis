import csv
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=13, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_note = Font(name=F, size=9, italic=True, color="666666")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_hdr = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF")
USD = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
PCT = '0.00%;(0.00%);-'
DT = 'yyyy-mm-dd'

wb = Workbook()

def hdr(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = f_bold; c.fill = fill_hdr
        c.alignment = Alignment(wrap_text=True, vertical="center")

def inp(c, v, fmt=None, key=False):
    c.value = v; c.font = f_in
    if key: c.fill = fill_in
    if fmt: c.number_format = fmt

def widths(ws, ws_w):
    for col, w in ws_w.items(): ws.column_dimensions[col].width = w

# ---------------- Guide ----------------
g = wb.active; g.title = "Guide"
lines = [
 ("NEST buyback model (Lido DAO), v1, 28.09.2026", f_title),
 ("", f_norm),
 ("How to use", f_bold),
 ("1. Variants sheet: each row is a separate buyback rule variant. Parameters are in the yellow cells (B–Z), results are calculated in the same row (AA–AN). Record your conclusion in column AO.", f_norm),
 ("2. Daily calculation of each variant on the Engine sheet, a separate block of columns for each variant.", f_norm),
 ("3. Stage1 and Stage2: precomputed sets of variants (values). Any row can be copied into Variants (columns B–Z match) to see the calculation with formulas.", f_norm),
 ("4. Forward: annual buyback capacity under constant conditions and a simplified 12-month forecast from the current state.", f_norm),
 ("", f_norm),
 ("Threshold type", f_bold),
 ("‘fixed’: amount per year; per day the integer part of amount / 365 is taken, as in the contract ($109,589 for $40M).", f_norm),
 ("‘expenses (actual)’: actual Foundations expenses for the period from reports. This is a retrospective estimate: the mechanism does not know the period's final expenses in real time. Use it as a reference for ‘what would have happened at this rate’.", f_norm),
 ("‘expenses (known)’: the latest published expense run rate as of that date (by report publication dates, table on the Expenses sheet). This is a feasible variant with no look-ahead.", f_norm),
 ("", f_norm),
 ("Profit measure (column Y) and evaluation metrics", f_bold),
 ("‘staking balance’ = net staking revenue − all Foundations expenses. ‘DAO operating’ = plus net Earn revenue (as Operating Result in the H1 2026 report). ‘cash flow’ = plus treasury income from own assets. ‘total with one-offs’ = minus one-off losses (Kelp, Q2 2026).", f_norm),
 ("‘Reduction is real’ (column Z): yes = the expense multiplier is also applied to the profit measure (expenses really were reduced), no = the threshold was lowered but expenses stayed the same.", f_norm),
 ("‘Buybacks in periods that ended in a loss’ (AL): purchases in reporting years or quarters whose result by the selected measure turned out negative. This is a hindsight estimate.", f_norm),
 ("‘Peak overshoot of buybacks over profit’ (AM): the maximum over the period of the difference between cumulative buybacks and cumulative DAO profit (from the start of the simulation). Shows by how much, at the worst moment, the mechanism bought more than the DAO had earned by then. No more than total buybacks.", f_norm),
 ("The number of LDO bought (AI) is illustrative: DefiLlama daily price, without slippage or execution delays. The main metric is buybacks in dollars.", f_norm),
 ("", f_norm),
 ("Colors", f_bold),
 ("Blue text on yellow background: inputs you can change. Blue text without background: source data (onchain or from reports). Black: formulas.", f_norm),
 ("", f_norm),
 ("What the revenue series is", f_bold),
 ("Actual staking fee inflows to the treasury (after the node operator share, before referral payouts), valued via Chainlink stETH/ETH × ETH/USD at the inflow block. This is the same oracle route as NEST uses, but not its historical counter: NEST converts to USD later, in a separate call. The stETH amount matches NEST exactly. The dollar valuation differs from the NEST counter by −0.044% in total over the live period, by individual day from −2.9% to +1.1%, 0.30% on average.", f_norm),
 ("Methodology details and all checks are in base/README.md.", f_norm),
 ("", f_norm),
 ("Assumptions", f_bold),
 ("Expenses for 2024, 2025 and H2 2026 by line item (staking, shared services) are unknown; for threshold variants 2 and 3 they are split using H1 2026 proportions. H2 2026 expenses = annual forecast $37.7M minus H1 actuals.", f_norm),
 ("The annual cap is computed in 365-day windows from the simulation start date, as the contract does from the activation date.", f_norm),
]
for i, (t, f) in enumerate(lines, 1):
    c = g.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
g.column_dimensions["A"].width = 130

# ---------------- Data ----------------
d = wb.create_sheet("Data")
rows = list(csv.DictReader(open("base/treasury_fee_inflows.csv")))
hdr(d, 1, ["Date (UTC)", "stETH to treasury", "stETH/ETH (Chainlink)", "ETH/USD (Chainlink)", "USD value", "Source",
           "Block", "Tx hash", "stETH/ETH roundId", "ETH/USD roundId", "Price flag"])
for i, r in enumerate(rows, 2):
    d.cell(row=i, column=1, value=date.fromisoformat(r["date"])).number_format = DT
    inp(d.cell(row=i, column=2), float(r["steth"]), "#,##0.000000")
    inp(d.cell(row=i, column=3), float(r["steth_eth"]), "0.000000")
    inp(d.cell(row=i, column=4), float(r["eth_usd"]), "#,##0.00")
    d.cell(row=i, column=5, value=f"=B{i}*C{i}*D{i}").number_format = USD2
    d.cell(row=i, column=6, value=r["source"])
    d.cell(row=i, column=7, value=int(r["block"]))
    d.cell(row=i, column=8, value=r["tx_hash"])
    d.cell(row=i, column=9, value=r["steth_eth_roundId"])
    d.cell(row=i, column=10, value=r["eth_usd_roundId"])
    d.cell(row=i, column=11, value=r["price_flag"])
    for col in (1, 5, 6, 7, 8, 9, 10, 11): d.cell(row=i, column=col).font = f_norm
N = len(rows); LAST = N + 1
d.freeze_panes = "A2"
widths(d, {"A": 12, "B": 16, "C": 14, "D": 14, "E": 14, "F": 24, "G": 11, "H": 68, "I": 22, "J": 22, "K": 26})
d.cell(row=1, column=2).comment = Comment("Actual stETH transfer to the treasury agent 0x3e40...9C8c: until 24.12.2025 minted at rebase, afterwards transferred from distributor 0x23ed...cddf. Source: base/treasury_fee_inflows.csv", "model")
d.cell(row=1, column=3).comment = Comment("Chainlink stETH/ETH proxy 0x86392dC19c0b719886221c78AB11eb8Cf5c52812, latestRoundData() at the inflow block", "model")
d.cell(row=1, column=4).comment = Comment("Chainlink ETH/USD proxy 0x5f4eC3Df9cbd43714FE2740f5E3616155c5b8419, latestRoundData() at the inflow block", "model")

DA = f"Data!$A$2:$A${LAST}"; DB = f"Data!$B$2:$B${LAST}"; DD = f"Data!$D$2:$D${LAST}"; DE = f"Data!$E$2:$E${LAST}"

# ---------------- Expenses ----------------
e = wb.create_sheet("Expenses")
e["A1"] = "Foundations expenses and cost of revenue by period"; e["A1"].font = f_title
hdr(e, 4, ["Period", "Start", "End", "Days", "Foundations expenses, $", "Staking (core+growth), $", "Shared services, $",
           "Onchain inflows, $", "Cost of revenue, $", "Cost share", "Daily threshold: all expenses",
           "Daily threshold: staking+shared", "Daily threshold: staking", "Report: net staking revenue, $", "Source / assumption"])
per = [
 ("2024", date(2024,1,1), date(2024,12,31)),
 ("2025", date(2025,1,1), date(2025,12,31)),
 ("Q1 2026", date(2026,1,1), date(2026,3,31)),
 ("Q2 2026", date(2026,4,1), date(2026,6,30)),
 ("H2 2026", date(2026,7,1), date(2026,12,31)),
]
for i, (lab, s, en) in enumerate(per, 5):
    e.cell(row=i, column=1, value=lab).font = f_norm
    e.cell(row=i, column=2, value=s).number_format = DT
    e.cell(row=i, column=3, value=en).number_format = DT
    e.cell(row=i, column=4, value=f"=C{i}-B{i}+1")
    e.cell(row=i, column=8, value=f'=SUMIFS({DE},{DA},">="&B{i},{DA},"<="&C{i})').number_format = USD
    e.cell(row=i, column=10, value=f"=IF(H{i}>0,I{i}/H{i},0)").number_format = PCT
    e.cell(row=i, column=11, value=f"=E{i}/D{i}*1").number_format = USD
    e.cell(row=i, column=12, value=f"=(F{i}+G{i})/D{i}").number_format = USD
    e.cell(row=i, column=13, value=f"=F{i}/D{i}").number_format = USD
# expenses totals
inp(e["E5"], 52100000, USD); inp(e["E6"], 45500000, USD); inp(e["E7"], 6920000, USD); inp(e["E8"], 7410000, USD)
e["E9"] = "=Expenses!$B$13-E7-E8"; e["E9"].number_format = USD
inp(e["F7"], 4380000, USD); inp(e["F8"], 4600000, USD); inp(e["G7"], 1570000, USD); inp(e["G8"], 1580000, USD)
for r in (5, 6, 9):
    e[f"F{r}"] = f"=E{r}*$B$14"; e[f"G{r}"] = f"=E{r}*$B$15"
    e[f"F{r}"].number_format = USD; e[f"G{r}"].number_format = USD
# cost of revenue
inp(e["I5"], 3100000, USD); inp(e["I6"], 4100000, USD)
inp(e["N7"], 8630000, USD); inp(e["N8"], 7080000, USD)
e["I7"] = "=H7-N7"; e["I8"] = "=H8-N8"; e["I9"] = "=H9*$B$16"
for c in ("I7", "I8", "I9"): e[c].number_format = USD
e["O5"] = "Expenses and referral payouts: GOOSE-2025 & EGGs-2025 Final Report (research.lido.fi/t/11304)"
e["O6"] = "Same"
e["O7"] = "H1 2026 Report, P&L (lido.fi/ldo-hub/reports/h1-2026); costs = onchain − net revenue"
e["O8"] = "Same"
e["O9"] = "Expenses = annual forecast (B13) − H1 actuals; line items and cost share by H1 proportions (assumption)"
for r in range(5, 10): e[f"O{r}"].font = f_note
e["A12"] = "Inputs and derived values"; e["A12"].font = f_bold
e["A13"] = "Expense forecast for 2026, $"; inp(e["B13"], 37700000, USD, key=True)
e["C13"] = "H1 2026 Report: full-year spending projected at ~$37.7M"; e["C13"].font = f_note
e["A14"] = "Staking share of expenses (H1 2026)"; e["B14"] = "=(F7+F8)/(E7+E8)"; e["B14"].number_format = PCT
e["A15"] = "Shared services share of expenses (H1 2026)"; e["B15"] = "=(G7+G8)/(E7+E8)"; e["B15"].number_format = PCT
e["A16"] = "Cost of revenue share (H1 2026)"; e["B16"] = "=(I7+I8)/(H7+H8)"; e["B16"].number_format = PCT
e["A17"] = "Onchain H2 2026 is incomplete (data up to the date of the last Data row); the H2 cost share is taken from H1."; e["A17"].font = f_note
widths(e, {"A": 40, "B": 14, "C": 12, "D": 7, "E": 17, "F": 17, "G": 15, "H": 17, "I": 17, "J": 10, "K": 14, "L": 14, "M": 14, "N": 17, "O": 80})


# ---- extra revenue lines and one-offs for profit measures ----
for col, t in (("Q", "Earn net revenue, $"), ("R", "Treasury income from own assets, $"), ("S", "One-off losses, $"),
               ("T", "Per day: staking balance"), ("U", "Per day: + Earn"), ("V", "Per day: + Earn + treasury income"), ("W", "Per day: + Earn + treasury income − one-offs")):
    e[f"{col}4"] = t; e[f"{col}4"].font = f_bold; e[f"{col}4"].fill = fill_hdr; e[f"{col}4"].alignment = Alignment(wrap_text=True, vertical="center")
inp(e["Q5"], 0, USD); inp(e["Q6"], 300000, USD); inp(e["Q7"], 200000, USD); inp(e["Q8"], 30000, USD)
e["Q9"] = "=(Q7+Q8)/(D7+D8)*D9"
inp(e["R5"], 3900000, USD); inp(e["R6"], 2700000, USD)
e["B18"] = "Treasury income H1 2026, $"; inp(e["C18"], 1400000, USD)
e["R7"] = "=$C$18*D7/(D7+D8)"; e["R8"] = "=$C$18*D8/(D7+D8)"; e["R9"] = "=$C$18/(D7+D8)*D9"
inp(e["S5"], 0, USD); inp(e["S6"], 0, USD); inp(e["S7"], 0, USD); inp(e["S8"], 6060000, USD); inp(e["S9"], 0, USD)
for r in range(5, 10):
    e[f"T{r}"] = 0; e[f"U{r}"] = f"=Q{r}/D{r}"; e[f"V{r}"] = f"=(Q{r}+R{r})/D{r}"; e[f"W{r}"] = f"=(Q{r}+R{r}-S{r})/D{r}"
    for c in "QRSTUVW": e[f"{c}{r}"].number_format = USD
for c, w in (("Q", 14), ("R", 16), ("S", 13), ("T", 12), ("U", 12), ("V", 14), ("W", 16)): e.column_dimensions[c].width = w
e["A19"] = ("Earn and treasury income: 2024–2025 from the GOOSE-2025 Final Report (Treasury Management, Lido Earn); H1 2026 from the H1 report (Net Lido Earn Revenue 0.20/0.03; treasury income $1.4M for the half-year, "
            "spread across days); H2 2026 at the H1 run rate (assumption). One-offs: Kelp −$6.06M in Q2 2026.")
e["A19"].font = f_note
# ---- known expense rates (no hindsight) ----
e["A21"] = "Expense run rate known as of the date (for the ‘expenses (known)’ threshold)"; e["A21"].font = f_bold
for i, t in enumerate(["From date", "All expenses, $/day", "Staking + shared, $/day", "Staking, $/day", "Source / assumption"]):
    c = e.cell(row=22, column=1 + i, value=t); c.font = f_bold; c.fill = fill_hdr
known = [
 (date(2024,1,1), "=52100000/366", "=52100000*($B$14+$B$15)/366", "=52100000*$B$14/366", "2024 actual; no data for 2023, for 2024 this is a look-ahead assumption"),
 (date(2025,1,1), "=52100000/366", "=52100000*($B$14+$B$15)/366", "=52100000*$B$14/366", "2024 run rate; publication date of the 2024 report not found, assumption: known from 01.01.2025"),
 (date(2026,3,17), "=45500000/365", "=45500000*($B$14+$B$15)/365", "=45500000*$B$14/365", "2025 run rate, report research.lido.fi/t/11304 dated 17.03.2026"),
 (date(2026,6,4), "=6920000/90", "=(4380000+1570000)/90", "=4380000/90", "Q1 2026 run rate, report research.lido.fi/t/11623 dated 04.06.2026"),
 (date(2026,8,28), "=14330000/181", "=(8980000+3150000)/181", "=8980000/181", "H1 2026 run rate, report research.lido.fi/t/11836 dated 28.08.2026"),
]
for i, (dt, a1, a2, a3, note) in enumerate(known, 23):
    e[f"A{i}"] = dt; e[f"A{i}"].number_format = DT
    for c, f in (("B", a1), ("C", a2), ("D", a3)):
        e[f"{c}{i}"] = f; e[f"{c}{i}"].number_format = USD
    e[f"E{i}"] = note; e[f"E{i}"].font = f_note
from openpyxl.utils import get_column_letter as CL
import json as _json

ldo = _json.load(open("data/ldo_px.json"))
d.cell(row=1, column=12, value="LDO/USD (DefiLlama, ffill ≤1 day)").font = f_bold
d.cell(row=1, column=12).fill = fill_hdr
for i, pxv in enumerate(ldo, 2):
    inp(d.cell(row=i, column=12), float(pxv), "0.0000")
d.column_dimensions["L"].width = 16
e["P4"] = "Operating staking balance for the period (based on available data), $"; e["P4"].font = f_bold; e["P4"].fill = fill_hdr
e["P4"].alignment = Alignment(wrap_text=True, vertical="center")
for r in range(5, 10):
    e[f"P{r}"] = f'=H{r}-I{r}-K{r}*COUNTIFS({DA},">="&B{r},{DA},"<="&C{r})'; e[f"P{r}"].number_format = USD
e.column_dimensions["P"].width = 20

# ---------------- Lists ----------------
ls = wb.create_sheet("Lists")
opts = {"A": ["fixed", "expenses (actual)", "expenses (known)"], "B": ["all", "staking+shared", "staking"], "C": ["as NEST", "net"],
        "D": ["accumulates", "floor N days", "does not accumulate", "decay", "reset by period", "rolling window"], "E": ["no", "progressive"],
        "F": ["quarter", "year"], "G": ["staking balance", "DAO operating", "cash flow", "summary incl. one-offs"], "H": ["yes", "no"]}
heads = {"A": "Threshold type", "B": "Expenses for threshold", "C": "Revenue base", "D": "Debt policy", "E": "Scale",
         "F": "Reset period", "G": "Profit measure", "H": "Reduction is real"}
for col, vals in opts.items():
    ls[f"{col}1"] = heads[col]; ls[f"{col}1"].font = f_bold
    for i, val in enumerate(vals, 2): ls[f"{col}{i}"] = val
ls["J1"] = "Helper sheet with lists for dropdown menus"; ls["J1"].font = f_note

# ---------------- Variants ----------------
v = wb.create_sheet("Variants", 1)
v["A1"] = "Buyback model variants: one row = one variant, results are computed in the same row"; v["A1"].font = f_title
v["A2"] = ("Change the yellow cells (B–Z). Results (AA–AN) recalculate automatically. Record your conclusion in column AO. "
           "Daily calculation of each variant on the Engine sheet. Parameter explanations below the table and on the Guide sheet."); v["A2"].font = f_note
cols = ["№", "Variant name", "Start", "End", "Threshold type", "Fixed threshold, $/year", "Expenses for threshold",
        "Expense multiplier", "Revenue base", "Share (or tier 1 share)", "Daily cap, $", "Annual cap, $", "Min. purchase, $",
        "Debt policy", "N days (floor / window)", "Initial budget, $",
        "Scale", "Bound 1, $/year above threshold", "Bound 2, $/year above threshold", "Tier 2 share", "Tier 3 share", "Share in deficit",
        "Deficit decay, % per day", "Reset period", "Profit measure for evaluation", "Expense reduction is real",
        "Buybacks 2024, $", "Buybacks 2025, $", "Buybacks 2026, $", "Total buybacks, $", "Days with purchases", "First purchase",
        "Budget at end, $", "Deepest deficit, $", "LDO bought, M (illustr.)", "DAO profit for the period (by measure Y), $",
        "Profit coverage", "Buybacks in periods that ended in a loss, $", "Peak overshoot of buybacks over profit, $", "Buyback activations",
        "Conclusion / comment", "",
        "aux: type", "helper: expense variant", "helper: as NEST?", "aux: debt code", "helper: scale?", "helper: quarter?",
        "aux: measure", "aux: profit multiplier",
        "aux: profit p1", "aux: profit p2", "aux: profit p3", "aux: profit p4", "aux: profit p5",
        "aux: buybacks p1", "aux: buybacks p2", "aux: buybacks p3", "aux: buybacks p4", "aux: buybacks p5"]
hdr(v, 4, cols)
v["B3"] = "Parameters"; v["B3"].font = f_bold
v["Q3"] = "Dynamic scale"; v["Q3"].font = f_bold
v["W3"] = "Debt"; v["W3"].font = f_bold
v["Y3"] = "Evaluation"; v["Y3"].font = f_bold
v["AA3"] = "Results"; v["AA3"].font = f_bold
v["AI3"] = "Comparison metrics"; v["AI3"].font = f_bold
S0, S1 = date(2025, 1, 1), date(2026, 9, 28)
D0 = dict(tp="fixed", fb=40e6, eo="all", em=1, rb="as NEST", sh=0.5, dm="accumulates", nd=30, sc="no", b1=5e6, b2=15e6,
          s2=0.5, s3=0.5, sn=0.5, dc=0, rp="quarter", ms="cash flow", cr="yes")
def V_(name, **kw):
    x = dict(D0); x.update(kw); x["name"] = name; return x
PROF = dict(tp="expenses (known)", rb="net", sh=1.0, s2=1.0, s3=1.0, sn=1.0)
variants = [
 V_("Current NEST"),
 V_("Current + deficit does not accumulate", dm="does not accumulate"),
 V_("Aksusarya proposal: 30M, 100%", fb=30e6, sh=1.0, s2=1.0, s3=1.0, sn=1.0),
 V_("Threshold = all expenses (actual), net, 100%, accumulates", **dict(PROF, tp="expenses (actual)")),
 V_("Threshold = all expenses (actual), reset by quarter", **dict(PROF, tp="expenses (actual)", dm="reset by period")),
 V_("Threshold = all expenses (known), accumulates", **PROF),
 V_("Threshold = all expenses (known), reset by quarter", **dict(PROF, dm="reset by period")),
 V_("Same with expenses ×0.8, reduction is real", **dict(PROF, dm="reset by period", em=0.8)),
 V_("Same with expenses ×0.8, reduction did NOT happen", **dict(PROF, dm="reset by period", em=0.8, cr="no")),
 V_("Threshold = staking+shared (known), reset by quarter", **dict(PROF, eo="staking+shared", dm="reset by period")),
 V_("Threshold = staking (known), 50%, reset by quarter", **dict(PROF, eo="staking", sh=0.5, s2=0.5, s3=0.5, sn=0.5, dm="reset by period")),
 V_("Scale at 40M threshold: 25% / 50% / 75%", sc="progressive", sh=0.25, s2=0.5, s3=0.75, sn=0.5),
 V_("Tier scale on expenses (known): 50% / 75% / 100%", **dict(PROF, sc="progressive", sh=0.5, s2=0.75, s3=1.0, sn=0.5, dm="reset by period")),
 V_("All expenses (known): decay 1% per day", **dict(PROF, dm="decay", dc=0.01)),
 V_("All expenses (known): window 90 days", **dict(PROF, dm="rolling window", nd=90)),
 V_("All expenses (known): reset by year", **dict(PROF, dm="reset by period", rp="year")),
 V_("Custom variant 1"),
 V_("Custom variant 2"),
]
NV = len(variants)
FIRST, LASTV = 5, 4 + NV
for k, x in enumerate(variants):
    r = FIRST + k
    v[f"A{r}"] = k + 1
    vals = [("B", x["name"], None), ("C", S0, DT), ("D", S1, DT), ("E", x["tp"], None), ("F", x["fb"], USD), ("G", x["eo"], None),
            ("H", x["em"], "0.00"), ("I", x["rb"], None), ("J", x["sh"], "0%"), ("K", 50000, USD), ("L", 10000000, USD),
            ("M", 1000, USD), ("N", x["dm"], None), ("O", x["nd"], "0"), ("P", 0, USD),
            ("Q", x["sc"], None), ("R", x["b1"], USD), ("S", x["b2"], USD), ("T", x["s2"], "0%"), ("U", x["s3"], "0%"), ("V", x["sn"], "0%"),
            ("W", x["dc"], "0.00%"), ("X", x["rp"], None), ("Y", x["ms"], None), ("Z", x["cr"], None)]
    for col, val, fmt in vals: inp(v[f"{col}{r}"], val, fmt, key=True)
    v[f"AQ{r}"] = f"=MATCH(E{r},Lists!$A$2:$A$4,0)"
    v[f"AR{r}"] = f"=MATCH(G{r},Lists!$B$2:$B$4,0)"
    v[f"AS{r}"] = f'=IF(I{r}="as NEST",1,0)'
    v[f"AT{r}"] = f"=MATCH(N{r},Lists!$D$2:$D$7,0)"
    v[f"AU{r}"] = f'=IF(Q{r}="progressive",1,0)'
    v[f"AV{r}"] = f'=IF(X{r}="quarter",1,0)'
    v[f"AW{r}"] = f"=MATCH(Y{r},Lists!$G$2:$G$5,0)"
    v[f"AX{r}"] = f'=IF(Z{r}="yes",H{r},1)'
for ref, rng in [("E", "$A$2:$A$4"), ("G", "$B$2:$B$4"), ("I", "$C$2:$C$3"), ("N", "$D$2:$D$7"), ("Q", "$E$2:$E$3"),
                 ("X", "$F$2:$F$3"), ("Y", "$G$2:$G$5"), ("Z", "$H$2:$H$3")]:
    dv = DataValidation(type="list", formula1=f"=Lists!{rng}", allow_blank=False); v.add_data_validation(dv)
    dv.add(f"{ref}{FIRST}:{ref}{LASTV}")
v.freeze_panes = "C5"
wd = {"A": 4, "B": 44, "C": 11, "D": 11, "E": 16, "F": 13, "G": 15, "H": 10, "I": 11, "J": 10, "K": 11, "L": 13, "M": 10,
      "N": 15, "O": 8, "P": 13, "Q": 13, "R": 13, "S": 13, "T": 9, "U": 9, "V": 9, "W": 10, "X": 10, "Y": 16, "Z": 11,
      "AA": 13, "AB": 13, "AC": 13, "AD": 14, "AE": 9, "AF": 12, "AG": 14, "AH": 14, "AI": 10, "AJ": 16, "AK": 11,
      "AL": 16, "AM": 16, "AN": 10, "AO": 50, "AP": 2}
for j in range(43, 61): wd[CL(j)] = 9
widths(v, wd)
v.row_dimensions[4].height = 72
note_r = LASTV + 2
v[f"A{note_r}"] = "How to read the parameters"; v[f"A{note_r}"].font = f_bold
notes = [
 "Threshold type: ‘fixed’ = amount per year (F), per day the integer part of F/365, as in the contract. ‘expenses (actual)’ = actual period expenses from reports, a hindsight estimate. ‘expenses (known)’ = latest published expense run rate as of the date, feasible with no look-ahead.",
 "Expenses for threshold: all Foundations expenses, staking together with shared services, or staking only. Multiplier 0.8 = threshold calculated for expenses 20% lower.",
 "Revenue base: ‘as NEST’ = treasury inflows before referral payouts, ‘net’ = after cost of revenue (as in Lido reports).",
 "Debt policy: accumulates (as now), floor N days, does not accumulate, decay W% per day, reset by period (quarter or year, column X), rolling window N days.",
 "‘Progressive’ tier scale: surplus above the threshold is split into tiers. Up to boundary 1 the share is J, between the boundaries the share is T, above boundary 2 the share is U. The deficit accumulates at share V.",
 "The profit measure (Y) sets which DAO profit the buybacks in columns AJ, AK, AL, AM are compared against. ‘Reduction is real’ (Z): yes = the expense multiplier is also applied to the profit measure.",
 "AL = buybacks in reporting years or quarters that ended in a loss by measure Y (a hindsight estimate). AM = by how much, at the worst moment of the period, cumulative buybacks exceeded cumulative profit (if profit is negative, all buybacks up to that moment count as uncovered). The value is no more than total buybacks.",
 "Caps and minimum purchase default to those in the NEST contract. 0 in a cap column = no cap. LDO bought is illustrative; the main metric is buybacks in dollars.",
]
for i, t in enumerate(notes, note_r + 1):
    v[f"A{i}"] = t; v[f"A{i}"].font = f_norm

# ---------------- Engine ----------------
en = wb.create_sheet("Engine")
E0 = 3; ELAST = E0 + N - 1
en["A1"] = "Daily calculation"; en["A1"].font = f_bold
hdr(en, 2, ["Date", "Year", "Period #", "Treasury inflow, $", "Net staking revenue/day, $", "LDO/USD", "Quarter key", "Known-rate #"])
for i in range(E0, ELAST + 1):
    dr = i - 1
    en[f"A{i}"] = f"=Data!A{dr}"; en[f"A{i}"].number_format = DT
    en[f"B{i}"] = f"=YEAR(A{i})"
    en[f"C{i}"] = f"=MATCH(A{i},Expenses!$B$5:$B$9,1)"
    en[f"D{i}"] = f"=Data!E{dr}"; en[f"D{i}"].number_format = USD
    en[f"E{i}"] = f"=D{i}*(1-INDEX(Expenses!$J$5:$J$9,C{i}))"; en[f"E{i}"].number_format = USD
    en[f"F{i}"] = f"=Data!L{dr}"; en[f"F{i}"].number_format = "0.0000"
    en[f"G{i}"] = f"=B{i}*10+ROUNDUP(MONTH(A{i})/3,0)"
    en[f"H{i}"] = f"=MATCH(A{i},Expenses!$A$23:$A$27,1)"
BLOCK = ["Active", "Revenue used", "Baseline/day", "Surplus", "Delta", "Budget pre-policy", "Budget after policy",
         "Cap window", "Spent in window before", "Buy", "Budget end", "Cum delta", "Cum buy",
         "DAO profit/day (measure)", "Cum profit in window", "Buys ahead of profit"]
C0 = 10
for k in range(NV):
    vr = FIRST + k
    c0 = C0 + k * (len(BLOCK) + 1)
    L = {nm: CL(c0 + j) for j, nm in enumerate(BLOCK)}
    en.cell(row=1, column=c0, value=f"=\"Variant {k+1}: \"&Variants!$B${vr}").font = f_bold
    hdr(en, 2, BLOCK, col=c0)
    V = lambda col: f"Variants!${col}${vr}"
    a, rv, bs, su, de, pre, fl, wi, sp, by, ed, cd, cb, pf, cp, uc = [L[n] for n in BLOCK]
    CDR = f"${cd}${E0}:${cd}${ELAST}"; CBR = f"${cb}${E0}:${cb}${ELAST}"
    for i in range(E0, ELAST + 1):
        p = i - 1
        en[f"{a}{i}"] = f"=IF(AND($A{i}>={V('C')},$A{i}<={V('D')}),1,0)"
        en[f"{rv}{i}"] = f"=IF({V('AS')}=1,$D{i},$E{i})"
        en[f"{bs}{i}"] = (f"=CHOOSE({V('AQ')},INT({V('F')}/365),INDEX(Expenses!$K$5:$M$9,$C{i},{V('AR')})*{V('H')},"
                          f"INDEX(Expenses!$B$23:$D$27,$H{i},{V('AR')})*{V('H')})")
        en[f"{su}{i}"] = f"={rv}{i}-{bs}{i}"
        dyn = (f"IF({su}{i}>=0,{V('J')}*MIN({su}{i},{V('R')}/365)+{V('T')}*MIN(MAX({su}{i}-{V('R')}/365,0),"
               f"MAX({V('S')}-{V('R')},0)/365)+{V('U')}*MAX({su}{i}-{V('S')}/365,0),{V('V')}*{su}{i})")
        en[f"{de}{i}"] = f"=IF({a}{i}=1,IF({V('AU')}=1,{dyn},{V('J')}*{su}{i}),0)"
        if i == E0:
            prevadj = V('P')
        else:
            prev = f"IF({a}{p}=1,{ed}{p},{V('P')})"
            newp = f"IF({V('AV')}=1,$G{i}<>$G{p},$B{i}<>$B{p})"
            prevadj = f"IF(AND({V('AT')}=5,{newp}),MAX(0,{prev}),{prev})"
        en[f"{pre}{i}"] = f"=IF({a}{i}=1,{prevadj}+{de}{i},0)"
        negshare = f"IF({V('AU')}=1,{V('V')},{V('J')})"
        nd = V('O')
        cdn = f"IF({i}-{nd}<{E0},0,INDEX({CDR},{i}-{nd}-{E0}+1))"
        cbn = f"IF({i}-{nd}<{E0},0,INDEX({CBR},{i}-{nd}-{E0}+1))"
        cbp = "0" if i == E0 else f"{cb}{p}"
        window = f"({cd}{i}-{cdn})-({cbp}-{cbn})"
        en[f"{fl}{i}"] = (f"=IF({a}{i}=1,CHOOSE({V('AT')},{pre}{i},MAX({pre}{i},-{nd}*{bs}{i}*{negshare}),MAX(0,{pre}{i}),"
                          f"IF({pre}{i}<0,{pre}{i}*(1-{V('W')}),{pre}{i}),{pre}{i},{window}),0)")
        en[f"{wi}{i}"] = f"=IF({a}{i}=1,INT(($A{i}-{V('C')})/365),-1)"
        en[f"{sp}{i}"] = "=0" if i == E0 else f"=IF(AND({a}{i}=1,{wi}{i}={wi}{p}),{sp}{p}+{by}{p},0)"
        en[f"{by}{i}"] = (f"=IF(AND({a}{i}=1,{fl}{i}>{V('M')}),MAX(0,MIN({fl}{i},IF({V('K')}>0,{V('K')},1E+18),"
                          f"IF({V('L')}>0,{V('L')}-{sp}{i},1E+18))),0)")
        en[f"{ed}{i}"] = f"={fl}{i}-{by}{i}"
        en[f"{cd}{i}"] = f"={de}{i}" if i == E0 else f"={cd}{p}+{de}{i}"
        en[f"{cb}{i}"] = f"={by}{i}" if i == E0 else f"={cb}{p}+{by}{i}"
        en[f"{pf}{i}"] = f"=$E{i}+INDEX(Expenses!$T$5:$W$9,$C{i},{V('AW')})-INDEX(Expenses!$K$5:$K$9,$C{i})*{V('AX')}"
        cpprev = "0" if i == E0 else f"IF({a}{p}=1,{cp}{p},0)"
        en[f"{cp}{i}"] = f"=IF({a}{i}=1,{cpprev}+{pf}{i},0)"
        en[f"{uc}{i}"] = f"=IF({a}{i}=1,MAX(0,{cb}{i}-MAX({cp}{i},0)),0)"
        for col in (rv, bs, su, de, pre, fl, sp, by, ed, cd, cb, pf, cp, uc): en[f"{col}{i}"].number_format = USD
    for j in range(len(BLOCK)): en.column_dimensions[CL(c0 + j)].width = 13
    en.column_dimensions[CL(c0 + len(BLOCK))].width = 2
    EA = f"Engine!$A${E0}:$A${ELAST}"; EB = f"Engine!$B${E0}:$B${ELAST}"; EC = f"Engine!$C${E0}:$C${ELAST}"
    BY = f"Engine!${by}${E0}:${by}${ELAST}"; ED = f"Engine!${ed}${E0}:${ed}${ELAST}"
    ACT = f"Engine!${a}${E0}:${a}${ELAST}"; PF = f"Engine!${pf}${E0}:${pf}${ELAST}"; UC = f"Engine!${uc}${E0}:${uc}${ELAST}"
    for col, y in (("AA", 2024), ("AB", 2025), ("AC", 2026)):
        v[f"{col}{vr}"] = f"=SUMIFS({BY},{EB},{y})"; v[f"{col}{vr}"].number_format = USD
    v[f"AD{vr}"] = f"=SUM(AA{vr}:AC{vr})"; v[f"AD{vr}"].number_format = USD; v[f"AD{vr}"].font = f_bold
    v[f"AE{vr}"] = f'=COUNTIF({BY},">0")'
    v[f"AF{vr}"] = f'=IF(AE{vr}>0,_xlfn.MINIFS({EA},{BY},">0"),"no")'; v[f"AF{vr}"].number_format = DT
    v[f"AG{vr}"] = f"=INDEX({ED},MATCH(D{vr},{EA},0))"; v[f"AG{vr}"].number_format = USD
    v[f"AH{vr}"] = f"=MIN({ED})"; v[f"AH{vr}"].number_format = USD
    v[f"AI{vr}"] = f"=SUMPRODUCT({BY}/Engine!$F${E0}:$F${ELAST})/1000000"; v[f"AI{vr}"].number_format = "0.00"
    v[f"AJ{vr}"] = f"=SUMIFS({PF},{ACT},1)"; v[f"AJ{vr}"].number_format = USD
    v[f"AK{vr}"] = f'=IF(AJ{vr}>0,AD{vr}/AJ{vr},"no profit")'; v[f"AK{vr}"].number_format = "0%"
    for j in range(5):
        pc = CL(51 + j); bc = CL(56 + j)
        v[f"{pc}{vr}"] = f"=SUMIFS({PF},{EC},{j+1})"; v[f"{pc}{vr}"].number_format = USD
        v[f"{bc}{vr}"] = f"=SUMIFS({BY},{EC},{j+1})"; v[f"{bc}{vr}"].number_format = USD
    v[f"AL{vr}"] = f"=SUMPRODUCT(({CL(51)}{vr}:{CL(55)}{vr}<0)*{CL(56)}{vr}:{CL(60)}{vr})"; v[f"AL{vr}"].number_format = USD
    v[f"AM{vr}"] = f"=MAX({UC})"; v[f"AM{vr}"].number_format = USD
    v[f"AN{vr}"] = (f"=IF(Engine!{by}{E0}>0,1,0)+SUMPRODUCT((Engine!{by}{E0+1}:{by}{ELAST}>0)*"
                    f"(Engine!{by}{E0}:{by}{ELAST-1}=0))")
en.column_dimensions["A"].width = 11; en.column_dimensions["D"].width = 14; en.column_dimensions["E"].width = 14
en.freeze_panes = "I3"
# ---------------- NEST_check ----------------
n = wb.create_sheet("NEST_check")
n["A1"] = "Reproduction of the NEST onchain budget at checkpoints"; n["A1"].font = f_title
n["A2"] = "Contract formula: budget = share × (counter revenue − daily threshold × number of daily slots). Slots are counted from the activation day inclusive through the checkpoint day inclusive."; n["A2"].font = f_note
n["A4"] = "NEST activation (activationTS)"; inp(n["B4"], date(2026, 8, 10), DT)
n["A5"] = "Daily threshold, $ (reserveDailyRateUSD)"; inp(n["B5"], 109589, USD)
n["A6"] = "Share (surplusShareBP)"; inp(n["B6"], 0.5, "0%")
n["A7"] = "First rebase captured by the counter"; inp(n["B7"], date(2026, 8, 15), DT)
n["C4"] = "Source: BuybackAllocator 0xAA568141c051f2D1132b110f8391F18D48E8D889, StakingRevenueSource 0x6220212a33a87Ed7Cc386B67eB2c393974F28C38"; n["C4"].font = f_note
hdr(n, 9, ["Checkpoint date", "Block", "Tx hash", "NEST revenue counter, $", "NEST onchain budget, $", "Slots",
           "Budget by formula, $", "Error, $", "Our revenue for the same rebases, $", "Budget on our revenue, $", "Difference from onchain, %"])
cps = list(csv.DictReader(open("base/nest_checkpoints.csv")))
for i, c in enumerate(cps, 10):
    dt = date.fromisoformat(c["ts"][:10])
    n[f"A{i}"] = dt; n[f"A{i}"].number_format = DT
    n[f"B{i}"] = int(c["block"]); n[f"C{i}"] = c["tx"]
    inp(n[f"D{i}"], float(c["totalRevenueUSD"]), USD2); inp(n[f"E{i}"], float(c["budgetUSD"]), USD2)
    n[f"F{i}"] = f"=A{i}-$B$4+1"
    n[f"G{i}"] = f"=$B$6*(D{i}-$B$5*F{i})"; n[f"G{i}"].number_format = USD2
    n[f"H{i}"] = f"=G{i}-E{i}"; n[f"H{i}"].number_format = USD2
    n[f"I{i}"] = f'=SUMIFS({DE},{DA},">="&$B$7,{DA},"<"&A{i})'; n[f"I{i}"].number_format = USD2
    n[f"J{i}"] = f"=$B$6*(I{i}-$B$5*F{i})"; n[f"J{i}"].number_format = USD2
    n[f"K{i}"] = f"=IF(D{i}>0,I{i}/D{i}-1,0)"; n[f"K{i}"].number_format = PCT
lastn = 9 + len(cps)
n[f"A{lastn+2}"] = "Maximum formula error, $"; n[f"B{lastn+2}"] = f"=MAX(MAX(H10:H{lastn}),-MIN(H10:H{lastn}))"
n[f"B{lastn+2}"].number_format = USD2
n[f"A{lastn+3}"] = "Latest NEST budget, $"; n[f"B{lastn+3}"] = f"=E{lastn}"; n[f"B{lastn+3}"].number_format = USD2
widths(n, {"A": 36, "B": 12, "C": 68, "D": 20, "E": 20, "F": 8, "G": 20, "H": 12, "I": 22, "J": 22, "K": 12})
NEST_LAST_BUDGET = f"NEST_check!$B${lastn+3}"
cr = lastn + 5
n[f"A{cr}"] = "NEST starting budget: three separate scenarios (not additive)"; n[f"A{cr}"].font = f_bold
n[f"A{cr+1}"] = "Allocator activated 10.08.2026 13:01 UTC (after that day's rebase at 12:21). Revenue source connected by DG #13 on 14.08.2026 13:23 UTC (after that day's rebase). First counted rebase 15.08.2026."; n[f"A{cr+1}"].font = f_note
n[f"A{cr+2}"] = "1. As in the contract (history is correct per specification), $"; n[f"B{cr+2}"] = f"=B{lastn+3}"
n[f"A{cr+3}"] = "2. Revenue source would have worked right after activation: count rebases 11–14.08 (revenue of 10.08 was before activation), $"
n[f"B{cr+3}"] = f'=B{lastn+3}+$B$6*SUMIFS({DE},{DA},">="&DATE(2026,8,11),{DA},"<="&DATE(2026,8,14))'
n[f"C{cr+3}"] = f'=$B$6*SUMIFS({DE},{DA},">="&DATE(2026,8,11),{DA},"<="&DATE(2026,8,14))'
n[f"A{cr+4}"] = "3. Reserve accrues only from the first full day of operation (15.08): remove slots 10–14.08, $"
n[f"B{cr+4}"] = f"=B{lastn+3}+5*$B$5*$B$6"; n[f"C{cr+4}"] = "=5*$B$5*$B$6"
n[f"D{cr+2}"] = "Adjustment, $"; n[f"D{cr+2}"].font = f_bold
for rr in (cr+2, cr+3, cr+4):
    n[f"B{rr}"].number_format = USD2; n[f"C{rr}"].number_format = USD2
n[f"A{cr+6}"] = ("Accounting timing shift: the checkpoint (~00:35 UTC) accrues the reserve for the new day before that day's rebase (~12:20 UTC), so at the moment of the checkpoint the budget is lower by one daily slot "
                 "($54,794.50) relative to accounting by completed days. This is a property of discrete slots per the specification, not a loss of revenue: a day's revenue is counted by the next checkpoint.")
n[f"A{cr+6}"].font = f_note

# ---------------- Recon ----------------
rc = wb.create_sheet("Recon")
rc["A1"] = "True-up of onchain inflows against Lido reports"; rc["A1"].font = f_title
hdr(rc, 3, ["Period", "Start", "End", "Onchain treasury inflows, $", "Report: net staking revenue, $", "Difference, $",
            "Report: referral payouts, $", "Residual, $", "Report: total protocol fee, $", "Treasury share of fee", "Source"])
recon = [
 ("2024", date(2024,1,1), date(2024,12,31), 48500000, 3100000, None, "GOOSE-2025 Final Report (t/11304)"),
 ("2025", date(2025,1,1), date(2025,12,31), 37400000, 4100000, None, "GOOSE-2025 Final Report (t/11304)"),
 ("Q1 2026", date(2026,1,1), date(2026,3,31), 8630000, None, 14990000, "H1 2026 Report, P&L"),
 ("Q2 2026", date(2026,4,1), date(2026,6,30), 7080000, None, 12520000, "H1 2026 Report, P&L"),
 ("H1 2026", date(2026,1,1), date(2026,6,30), 15710000, None, 27510000, "H1 2026 Report, P&L"),
]
for i, (lab, a, b, net, ref, gross, src) in enumerate(recon, 4):
    rc[f"A{i}"] = lab; rc[f"B{i}"] = a; rc[f"C{i}"] = b
    rc[f"B{i}"].number_format = DT; rc[f"C{i}"].number_format = DT
    rc[f"D{i}"] = f'=SUMIFS({DE},{DA},">="&B{i},{DA},"<="&C{i})'
    inp(rc[f"E{i}"], net, USD)
    rc[f"F{i}"] = f"=D{i}-E{i}"
    if ref is not None:
        inp(rc[f"G{i}"], ref, USD); rc[f"H{i}"] = f"=F{i}-G{i}"
    else:
        rc[f"G{i}"] = "not disclosed"; rc[f"H{i}"] = "n/a"
    if gross is not None:
        inp(rc[f"I{i}"], gross, USD); rc[f"J{i}"] = f"=D{i}/I{i}"; rc[f"J{i}"].number_format = PCT
    rc[f"K{i}"] = src; rc[f"K{i}"].font = f_note
    for c in "DFH": rc[f"{c}{i}"].number_format = USD
rc["A10"] = "The 2024/2025 residual is within report rounding ($0.1M) and the difference in valuation methods (report at end-of-day price, onchain at block price)."; rc["A10"].font = f_note
rc["A11"] = "The H1 2026 difference is cost of revenue: under Lido's methodology, referral payouts and possibly operator discounts; there is no breakdown for H1."; rc["A11"].font = f_note
widths(rc, {"A": 10, "B": 12, "C": 12, "D": 22, "E": 22, "F": 14, "G": 18, "H": 12, "I": 20, "J": 12, "K": 36})

# ---------------- Forward ----------------
fw = wb.create_sheet("Forward")
fw["A1"] = "Annual buyback capacity under constant conditions and a simplified forecast: revenue = TVL × yield × DAO share × ETH price"; fw["A1"].font = f_title
fw["A3"] = "Calibration to actual data"; fw["A3"].font = f_bold
fw["A4"] = "Average treasury inflows over 30 days, stETH/day"
fw["B4"] = f'=AVERAGEIFS({DB},{DA},">"&(MAX({DA})-30))'; fw["B4"].number_format = "0.00"
fw["A5"] = "Latest ETH/USD price"; fw["B5"] = f"=INDEX({DD},MATCH(MAX({DA}),{DA},0))"; fw["B5"].number_format = "#,##0.00"
fw["A6"] = "DAO share of rewards"; inp(fw["B6"], 0.0617, PCT, key=True)
fw["C6"] = "StakingRouter getStakingFeeAggregateDistributionE4Precision(), 28.09.2026"; fw["C6"].font = f_note
fw["A7"] = "Lido TVL, M ETH"; inp(fw["B7"], 9.13, "0.00", key=True)
fw["C7"] = "H1 2026 Report, 30.06.2026, with entry queue; the implied yield below is calibrated so that revenue matches actuals"; fw["C7"].font = f_note
fw["A8"] = "Implied gross staking yield"; fw["B8"] = "=B4*365/(B6*B7*1000000)"; fw["B8"].number_format = PCT
fw["A9"] = "Current NEST budget, $"; fw["B9"] = f"={NEST_LAST_BUDGET}"; fw["B9"].number_format = USD; fw["B9"].font = f_link
fw["A10"] = "Annual cap, $"; inp(fw["B10"], 10000000, USD)
fw["A11"] = "Share for expense-based rule"; inp(fw["B11"], 1.0, "0%", key=True)

hdr(fw, 13, ["Metric", "Now", "Team targets (23% share)", "ETH $2,000", "ETH $3,500", "Issuance −30%", "Move to Curated v2 (5.8%)"])
cols = "BCDEFG"
fw["A14"] = "TVL, M ETH"; fw["A15"] = "Gross yield"; fw["A16"] = "DAO share"; fw["A17"] = "ETH price, $"
fw["A18"] = "Cost of revenue share"; fw["A19"] = "Foundations expenses, 2026 forecast, $ per year"; fw["A20"] = "NEST fixed threshold, $"
for c in cols:
    fw[f"{c}14"] = "=$B$7"; fw[f"{c}15"] = "=$B$8"; fw[f"{c}16"] = "=$B$6"; fw[f"{c}17"] = "=$B$5"
    fw[f"{c}18"] = "=Expenses!$B$16"; fw[f"{c}19"] = "=Expenses!$B$13"
    fw[f"{c}20"] = 40000000; fw[f"{c}20"].font = f_in
fw["C14"] = "=$B$7*23/21.2"
fw["D17"] = 2000; fw["D17"].font = f_in; fw["D17"].fill = fill_in
fw["E17"] = 3500; fw["E17"].font = f_in; fw["E17"].fill = fill_in
fw["F15"] = "=$B$8*0.7"
fw["G16"] = 0.058; fw["G16"].font = f_in; fw["G16"].fill = fill_in
fw["A22"] = "Results"; fw["A22"].font = f_bold
labels = {23: "Treasury inflows, $ per year", 24: "Net staking revenue, $ per year",
          25: "Break-even ETH price for the NEST threshold", 26: "ETH price at which net revenue = expenses",
          27: "Annual capacity: current rule (50%, cap), $/year", 28: "Days to pay off current deficit (current rule)",
          29: "Annual capacity: threshold = 2026 expense forecast, net, share B11, $/year", 30: "Annual capacity: proposal 30M / 100%, $/year",
          31: "Threshold based on expenses known as of the date (latest published run rate), $/year", 32: "Annual capacity: threshold = known expenses, net, share B11, $/year",
          33: "ETH price at which net revenue = known expense run rate", 34: "12-month forecast from the current state: current rule, $"}
for r, t in labels.items(): fw[f"A{r}"] = t
for c in cols:
    fw[f"{c}23"] = f"={c}14*1000000*{c}15*{c}16*{c}17"
    fw[f"{c}24"] = f"={c}23*(1-{c}18)"
    fw[f"{c}25"] = f"=INT({c}20/365)*365/({c}14*1000000*{c}15*{c}16)"
    fw[f"{c}26"] = f"={c}19/({c}14*1000000*{c}15*{c}16*(1-{c}18))"
    fw[f"{c}27"] = f"=MIN($B$10,0.5*MAX(0,{c}23-{c}20))"
    fw[f"{c}28"] = f'=IF({c}23>{c}20,-$B$9/(0.5*({c}23-{c}20)/365),"not paid off")'
    fw[f"{c}29"] = f"=MIN($B$10,$B$11*MAX(0,{c}24-{c}19))"
    fw[f"{c}30"] = f"=MIN($B$10,MAX(0,{c}23-30000000))"
    fw[f"{c}31"] = "=Expenses!$B$27*365"
    fw[f"{c}32"] = f"=MIN($B$10,$B$11*MAX(0,{c}24-{c}31))"
    fw[f"{c}33"] = f"={c}31/({c}14*1000000*{c}15*{c}16*(1-{c}18))"
    fw[f"{c}34"] = f"=MAX(0,MIN($B$10,0.5*MAX(0,{c}23-{c}20)+$B$9))"
    for r in (31, 32, 34): fw[f"{c}{r}"].number_format = USD
    fw[f"{c}33"].number_format = "$#,##0"
    for r in (19, 20, 23, 24, 27, 29, 30): fw[f"{c}{r}"].number_format = USD
    for r in (25, 26, 17): fw[f"{c}{r}"].number_format = "$#,##0"
    fw[f"{c}14"].number_format = "0.00"; fw[f"{c}15"].number_format = PCT; fw[f"{c}16"].number_format = PCT; fw[f"{c}18"].number_format = PCT
    fw[f"{c}28"].number_format = "#,##0"
fw["A36"] = ("‘Capacity’ rows are the annual buyback volume at constant price, TVL, yield and expenses, without volatility, with the budget starting from zero. "
             "This is not a forecast of actual buybacks. Row 34 accounts for the current NEST deficit (B9) at constant revenue, with no volatility and no remaining cap window.")
fw["A36"].font = f_note
fw["A37"] = ("The ‘known expenses’ threshold currently uses the H1 2026 run rate ($28.9M/year), while the team's forecast for the full year 2026 of $37.7M implies an H2 run rate of about $46M/year. "
             "A feasible rule without a quarterly true-up will buy earlier than the second-half expenses arrive.")
fw["A37"].font = f_note
fw["A39"] = "ETH break-even price: threshold level and revenue base separately"; fw["A39"].font = f_bold
for j, hname in enumerate(["Threshold", "$ per year", "Revenue as NEST", "Net revenue"]):
    c = fw.cell(row=40, column=1 + j, value=hname); c.font = f_bold; c.fill = fill_hdr
bk = [("Current threshold (contract)", "=INT(40000000/365)*365"), ("Expense forecast for 2026", "=Expenses!$B$13"),
      ("Known expense run rate (latest report)", "=Expenses!$B$27*365"), ("H2 2026 expense run rate per forecast", "=Expenses!$K$9*365")]
for i, (nm, f) in enumerate(bk, 41):
    fw[f"A{i}"] = nm; fw[f"B{i}"] = f; fw[f"B{i}"].number_format = USD
    fw[f"C{i}"] = f"=B{i}/($B$4*365)"; fw[f"D{i}"] = f"=B{i}/($B$4*365*(1-Expenses!$B$16))"
    fw[f"C{i}"].number_format = "$#,##0"; fw[f"D{i}"].number_format = "$#,##0"
fw["A46"] = ("Switching to an expense-based threshold changes two things at once: the threshold level and the revenue base. Lowering the level from $40M to $37.7M by itself lowers the break-even price, "
             "while switching to net revenue raises it by about 8%. They need to be assessed separately.")
fw["A46"].font = f_note
widths(fw, {"A": 58, "B": 16, "C": 18, "D": 14, "E": 14, "F": 14, "G": 20})


# ---------------- Stage1 / Stage2 (values from engine.py) ----------------
import os
VER = set(_json.load(open("/tmp/verified_v3.json"))) if os.path.exists("/tmp/verified_v3.json") else set()
PKEYS = ["start", "end", "type", "fixed", "exp", "mult", "basis", "share", "dcap", "acap", "minspend", "deficit", "ndays",
         "startbudget", "scale", "b1", "b2", "s2", "s3", "sneg", "decay", "reset", "measure", "cuts_real"]
PFMT = {"fixed": USD, "dcap": USD, "acap": USD, "minspend": USD, "startbudget": USD, "b1": USD, "b2": USD, "share": "0%",
        "s2": "0%", "s3": "0%", "sneg": "0%", "decay": "0.0%", "mult": "0.00"}
PHDR = ["№", "Variant", "Start", "End", "Threshold type", "Fixed threshold, $/year", "Expenses for threshold", "Expense multiplier", "Revenue base",
        "Share (tier 1)", "Daily cap, $", "Annual cap, $", "Min. purchase, $", "Debt policy", "N days", "Initial budget, $",
        "Scale", "Bound 1", "Bound 2", "Share 2", "Share 3", "Share in deficit", "Decay, %/day", "Reset period",
        "Profit measure", "Reduction is real"]
fill_a = PatternFill("solid", fgColor="F2F2F2"); fill_grp = PatternFill("solid", fgColor="FFF2CC")

def write_params(ws, row, p):
    ws.cell(row=row, column=2, value=p["name"]).font = f_bold
    for j, k in enumerate(PKEYS, 3):
        val = p[k]
        if k in ("start", "end"): val = date.fromisoformat(val)
        c = ws.cell(row=row, column=j, value=val); c.font = f_norm
        c.number_format = DT if k in ("start", "end") else PFMT.get(k, "General")

def header(ws, row, extra):
    for i, h in enumerate(PHDR + extra, 1):
        c = ws.cell(row=row, column=i, value=h); c.font = f_bold; c.fill = fill_hdr
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 72

def cov(r): return r["coverage"] if r["coverage"] is not None else "no profit"
def fdate(x): return date.fromisoformat(x) if x else "no"

s1 = wb.create_sheet("Stage1", 2)
s1["A1"] = "Stage 1: one change at a time from three bases"; s1["A1"].font = f_title
for i, t in enumerate([
    "Base A = current NEST. Base B = threshold equals all Foundations expenses per actual period data (hindsight estimate), net revenue, share 100%, deficit accumulates. "
    "Base C = the same, but threshold based on expenses known as of the date (feasible with no look-ahead).",
    "Each row changes exactly one parameter relative to its base, column AO shows the effect. Columns B–Z match the Variants sheet: copy a row there to see the calculation with formulas.",
    "Calculated by engine.py (replicates the spreadsheet formulas). Rows marked in AP were recalculated with live formulas and matched. Period 01.01.2025–28.09.2026 unless stated otherwise. Default profit measure is ‘cash flow’."], 2):
    s1[f"A{i}"] = t; s1[f"A{i}"].font = f_note
header(s1, 6, ["Buybacks 2024, $", "Buybacks 2025, $", "Buybacks 2026, $", "Total buybacks, $", "Days with purchases", "First purchase",
               "Budget at end, $", "Deepest deficit, $", "LDO bought, M (illustr.)", "DAO profit for the period, $", "Profit coverage",
               "Buybacks in periods that ended in a loss, $", "Peak overshoot of buybacks over profit, $", "Activations",
               "Effect of the change vs. base, $", "Verified with formulas"])
for n_, x in enumerate(_json.load(open("stage1_results.json")), 1):
    row = 6 + n_; r = x["res"]
    s1.cell(row=row, column=1, value=n_).font = f_norm
    write_params(s1, row, x["params"])
    vals = [(r["b2024"], USD), (r["b2025"], USD), (r["b2026"], USD), (r["total"], USD), (r["days"], "0"), (fdate(r["first"]), DT),
            (r["end"], USD), (r["min"], USD), (r["ldo_m"], "0.00"), (r["profit"], USD), (cov(r), "0%"), (r["over_period"], USD),
            (r["uncovered"], USD), (r["switches"], "0"), (x["delta"], '+$#,##0;-$#,##0;0'), ("yes" if x["params"]["name"] in VER else "", "@")]
    for j, (val, fmt) in enumerate(vals, 27):
        c = s1.cell(row=row, column=j, value=val); c.number_format = fmt; c.font = f_bold if j in (30, 41) else f_norm
    if x["tag"] == "A":
        for j in range(1, 43): s1.cell(row=row, column=j).fill = fill_a
for j, w in {1: 4, 2: 50}.items(): s1.column_dimensions[CL(j)].width = w
for j in range(3, 27): s1.column_dimensions[CL(j)].width = 11
for j in range(27, 43): s1.column_dimensions[CL(j)].width = 14
s1.freeze_panes = "C7"

s2 = wb.create_sheet("Stage2", 3)
s2["A1"] = "Stage 2: combinations"; s2["A1"].font = f_title
for i, t in enumerate([
    "Period 01.01.2025–28.09.2026. Columns B–Z match the Variants sheet. Calculated by engine.py; rows marked in the last column were recalculated with live formulas and matched.",
    "DAO profit by the ‘cash flow’ measure (net staking revenue + Earn + treasury income − all Foundations expenses), −$5.6M for the period at actual expenses.",
    "‘Reduction is real = no’ means the threshold was lowered but expenses stayed the same: compare such rows with their paired rows where the reduction happened.",
    "Annual capacity = buyback volume over a year under constant conditions starting from zero; this is not a forecast. The 12-month forecast (from 29.09.2026) is calculated with the same daily engine at a constant ETH price: "
    "for rules based on the current NEST with the current deficit of −$548,925 and a cap window from 10.08.2026, for new rules from zero. Forward expenses: H2 2026 forecast, from 2027 the 2026 forecast run rate (assumption)."], 2):
    s2[f"A{i}"] = t; s2[f"A{i}"].font = f_note
header(s2, 7, ["Total buybacks, $", "Days with purchases", "LDO bought, M (illustr.)", "Budget at end, $", "DAO profit for the period, $",
               "Profit coverage", "Buybacks in periods that ended in a loss, $", "Peak overshoot of buybacks over profit, $",
               "Forward threshold, $/year", "ETH break-even price", "Capacity, ETH $2,000", "Capacity, ETH $2,680", "Capacity, ETH $3,500",
               "Budget at forecast start, $", "12-month forecast, ETH $2,000", "12-month forecast, ETH $2,680", "12-month forecast, ETH $3,000", "12-month forecast, ETH $3,500",
               "Days to first purchase, ETH $2,680", "Days to first purchase, ETH $3,500",
               "Forecast ETH $3,000: peak overshoot over profit, $", "Forecast ETH $3,000: DAO profit, $", "Verified with formulas"])
row = 7; n_ = 0; last = None
for x in _json.load(open("stage2_results.json")):
    if x["group"] != last:
        row += 1; s2.cell(row=row, column=2, value=x["group"]).font = f_bold
        for j in range(1, 50): s2.cell(row=row, column=j).fill = fill_grp
        last = x["group"]
    row += 1; n_ += 1; r = x["res"]
    s2.cell(row=row, column=1, value=n_).font = f_norm
    write_params(s2, row, x["params"])
    d2680 = x["fc_days_to_first_2680"]; d3500 = x["fc_days_to_first_3500"]
    vals = [(r["total"], USD), (r["days"], "0"), (r["ldo_m"], "0.00"), (r["end"], USD), (r["profit"], USD), (cov(r), "0%"),
            (r["over_period"], USD), (r["uncovered"], USD), (x["base_annual_fwd"], USD), (x["breakeven"], "$#,##0"),
            (x["cap2000"], USD), (x["cap2680"], USD), (x["cap3500"], USD), (x["fc_start_budget"], USD),
            (x["fc2000"], USD), (x["fc2680"], USD), (x["fc3000"], USD), (x["fc3500"], USD),
            (d2680 if d2680 is not None else "no buybacks", "0"), (d3500 if d3500 is not None else "no buybacks", "0"),
            (x["fc_uncov_3000"], USD), (x["fc_profit_3000"], USD), ("yes" if x["params"]["name"] in VER else "", "@")]
    for j, (val, fmt) in enumerate(vals, 27):
        c = s2.cell(row=row, column=j, value=val); c.number_format = fmt; c.font = f_bold if j == 27 else f_norm
for j, w in {1: 4, 2: 58}.items(): s2.column_dimensions[CL(j)].width = w
for j in range(3, 27): s2.column_dimensions[CL(j)].width = 11
for j in range(27, 50): s2.column_dimensions[CL(j)].width = 14
s2.freeze_panes = "C8"
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != F:
                c.font = Font(name=F, size=10, bold=bool(c.font.bold), italic=bool(c.font.italic), color=(c.font.color.rgb if c.font.color is not None and isinstance(c.font.color.rgb, str) else None))

wb.save("../nest_model.xlsx")
print("saved")
